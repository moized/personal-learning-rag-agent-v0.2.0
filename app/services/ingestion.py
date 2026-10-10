from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse, parse_qs
from urllib.parse import quote
import ast
import json
import re

from pypdf import PdfReader


@dataclass
class RawDocument:
    external_id: str
    title: str
    section: str | None
    text: str
    metadata: dict | None = None


def clean_caption_text(text: str) -> str:
    text = re.sub(r"^WEBVTT.*$", "", text, flags=re.MULTILINE | re.IGNORECASE)
    text = re.sub(r"^NOTE.*(?:\n.*)*?(?=\n\n|\Z)", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\d+\s*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"\d{2}:\d{2}:\d{2}[,.]\d{3}\s*-->\s*\d{2}:\d{2}:\d{2}[,.]\d{3}.*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"\d{2}:\d{2}(?:\.\d{3})?\s*-->\s*\d{2}:\d{2}(?:\.\d{3})?.*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"<[^>]+>", "", text)
    return re.sub(r"\n{2,}", "\n", text).strip()


def parse_structured_literal(raw: str) -> dict:
    """Accept strict JSON or a Python-dict-like pasted object without executing arbitrary code."""
    raw = raw.strip()
    if not raw:
        raise ValueError("Structured input is empty")
    try:
        obj = json.loads(raw)
    except json.JSONDecodeError:
        try:
            obj = ast.literal_eval(raw)
        except (ValueError, SyntaxError) as exc:
            raise ValueError("Input is neither valid JSON nor a safe Python dictionary literal") from exc
    if not isinstance(obj, dict):
        raise ValueError("Structured input must be an object/dictionary")
    return obj


def _split_paragraphs(text: str) -> list[str]:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks = [re.sub(r"\s+", " ", b).strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    if blocks:
        return blocks
    return [re.sub(r"\s+", " ", text).strip()] if text.strip() else []


def _split_long(block: str, size: int) -> list[str]:
    if len(block) <= size:
        return [block]
    sentences = re.split(r"(?<=[.!?])\s+", block)
    if len(sentences) <= 1:
        words = block.split()
        out, current = [], []
        for word in words:
            candidate = (" ".join(current + [word])).strip()
            if current and len(candidate) > size:
                out.append(" ".join(current))
                current = [word]
            else:
                current.append(word)
        if current:
            out.append(" ".join(current))
        return out
    out, current = [], ""
    for sentence in sentences:
        if not sentence:
            continue
        candidate = f"{current} {sentence}".strip()
        if current and len(candidate) > size:
            out.append(current)
            current = sentence
        else:
            current = candidate
    if current:
        out.append(current)
    return out


def chunk_text(text: str, size: int = 1800, overlap: int = 220) -> list[str]:
    """Structure-aware chunking: paragraph/sentence boundaries first, then bounded overlap."""
    paragraphs = _split_paragraphs(text)
    pieces: list[str] = []
    for paragraph in paragraphs:
        pieces.extend(_split_long(paragraph, size))
    out: list[str] = []
    current = ""
    for piece in pieces:
        if not current:
            current = piece
            continue
        candidate = f"{current}\n{piece}"
        if len(candidate) <= size:
            current = candidate
        else:
            out.append(current.strip())
            tail = current[-overlap:].strip() if overlap else ""
            current = f"{tail}\n{piece}".strip() if tail else piece
            if len(current) > size:
                out.extend(_split_long(current, size)[:-1])
                current = _split_long(current, size)[-1]
    if current.strip():
        out.append(current.strip())
    return [x for x in out if x]


def parse_file(path: Path) -> list[RawDocument]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        docs = []
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                docs.append(RawDocument(f"{path.name}:p{i}", path.stem, f"page {i}", text))
        return docs
    if suffix in {".txt", ".md"}:
        return [RawDocument(path.name, path.stem, None, path.read_text(encoding="utf-8", errors="ignore"))]
    if suffix in {".srt", ".vtt"}:
        clean = clean_caption_text(path.read_text(encoding="utf-8", errors="ignore"))
        return [RawDocument(path.name, path.stem, "manual transcript", clean)]
    if suffix == ".json":
        payload = parse_structured_literal(path.read_text(encoding="utf-8", errors="ignore"))
        text = payload.get("transcript") or payload.get("text") or payload.get("content")
        if not isinstance(text, str) or not text.strip():
            raise ValueError("JSON/dict source must contain transcript, text, or content")
        title = str(payload.get("name") or payload.get("title") or path.stem)
        return [RawDocument(path.name, title, payload.get("section"), text, payload)]
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [RawDocument(path.name, path.stem, None, text)]
    if suffix == ".pptx":
        from pptx import Presentation
        presentation = Presentation(str(path))
        docs: list[RawDocument] = []
        for index, slide in enumerate(presentation.slides, start=1):
            pieces: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = shape.text.strip()
                    if value:
                        pieces.append(value)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            pieces.append(" | ".join(cells))
            text = "\n\n".join(pieces)
            if text.strip():
                docs.append(RawDocument(f"{path.name}:slide{index}", path.stem, f"slide {index}", text))
        return docs
    raise ValueError(f"Unsupported file type: {suffix}")


YOUTUBE_HOSTS = {
    "youtu.be", "youtube.com", "www.youtube.com", "m.youtube.com", "music.youtube.com",
    "youtube-nocookie.com", "www.youtube-nocookie.com",
}


def youtube_video_id(url: str) -> str | None:
    parsed = urlparse(url)
    host = (parsed.hostname or "").lower()
    if host == "youtu.be":
        video_id = parsed.path.strip("/").split("/")[0]
        return video_id if re.fullmatch(r"[A-Za-z0-9_-]{11}", video_id) else None
    if host in YOUTUBE_HOSTS:
        q = parse_qs(parsed.query).get("v")
        if q:
            return q[0] if re.fullmatch(r"[A-Za-z0-9_-]{11}", q[0]) else None
        m = re.search(r"/(?:shorts|embed)/([^/?]+)", parsed.path)
        return m.group(1) if m and re.fullmatch(r"[A-Za-z0-9_-]{11}", m.group(1)) else None
    return None


def youtube_is_playlist(url: str) -> bool:
    return bool(parse_qs(urlparse(url).query).get("list"))


def canonical_youtube_url(url: str) -> str:
    """Remove video tracking parameters so equivalent YouTube links deduplicate."""
    parsed = urlparse(url)
    if (parsed.hostname or "").lower() not in YOUTUBE_HOSTS:
        return url
    playlist_id = parse_qs(parsed.query).get("list", [None])[0]
    if playlist_id:
        return f"https://www.youtube.com/playlist?list={quote(playlist_id, safe='')}"
    video_id = youtube_video_id(url)
    if video_id:
        return f"https://www.youtube.com/watch?v={video_id}"
    return url


def _iso_date(value) -> str | None:
    """yt-dlp dates look like 20240115."""
    value = str(value or "")
    return f"{value[:4]}-{value[4:6]}-{value[6:8]}" if re.fullmatch(r"\d{8}", value) else None


def _video_entry(info: dict) -> dict:
    duration = info.get("duration")
    return {
        "id": info["id"],
        "title": info.get("title") or info["id"],
        "channel": info.get("channel") or info.get("uploader") or "",
        "upload_date": _iso_date(info.get("upload_date") or info.get("release_date")),
        "duration_seconds": int(duration) if isinstance(duration, (int, float)) else None,
    }


def list_youtube_entries(url: str) -> tuple[dict, list[dict]]:
    """Return (source info, videos) for a YouTube video or playlist.

    Titles, channel and upload date are included when YouTube provides them. For a
    single video the metadata lookup is best-effort: its transcript can still be
    fetched without it, so a metadata failure falls back to the video ID.
    """
    if (urlparse(url).hostname or "").lower() not in YOUTUBE_HOSTS:
        raise ValueError("URL must be a YouTube video or playlist URL")
    if not youtube_is_playlist(url):
        vid = youtube_video_id(url)
        if not vid:
            raise ValueError("Could not parse YouTube video URL")
        try:
            from yt_dlp import YoutubeDL
            options = {"quiet": True, "no_warnings": True, "skip_download": True, "ignore_no_formats_error": True}
            with YoutubeDL(options) as ydl:
                info = ydl.extract_info(f"https://www.youtube.com/watch?v={vid}", download=False) or {}
            entry = _video_entry({**info, "id": vid})
        except Exception:
            entry = {"id": vid, "title": vid, "channel": "", "upload_date": None, "duration_seconds": None}
        title = entry["title"] if entry["title"] != vid else ""
        return {"kind": "video", "title": title, "channel": entry["channel"], "upload_date": entry["upload_date"]}, [entry]
    try:
        from yt_dlp import YoutubeDL
        with YoutubeDL({"quiet": True, "extract_flat": True, "skip_download": True}) as ydl:
            info = ydl.extract_info(url, download=False) or {}
        videos = [_video_entry(entry) for entry in info.get("entries", []) if entry and entry.get("id")]
        if not videos:
            raise ValueError("Playlist contains no accessible videos")
    except Exception as exc:
        raise RuntimeError(f"Playlist expansion failed: {exc}") from exc
    source_info = {"kind": "playlist", "title": info.get("title") or "", "channel": info.get("channel") or info.get("uploader") or "",
                   "upload_date": _iso_date(info.get("modified_date"))}
    return source_info, videos


def list_youtube_videos(url: str) -> list[tuple[str, str]]:
    """(video_id, title) pairs for a YouTube video or playlist."""
    return [(video["id"], video["title"]) for video in list_youtube_entries(url)[1]]


class TranscriptUnavailable(RuntimeError):
    """The video itself has no usable transcript. Retrying will not help; the user can paste one."""

    def __init__(self, message: str, user_message: str):
        super().__init__(message)
        self.user_message = user_message


# youtube-transcript-api errors that describe the video, not a temporary block such as
# IpBlocked or RequestBlocked. Matched by name so the private error module is not imported.
_UNAVAILABLE_TRANSCRIPT_MESSAGES = {
    "NoTranscriptFound": "This video has no captions.",
    "TranscriptsDisabled": "Captions are turned off for this video.",
    "VideoUnavailable": "The video is unavailable (private, removed or region-blocked).",
    "VideoUnplayable": "The video cannot be played, so its captions cannot be read.",
    "InvalidVideoId": "The video ID is not valid.",
    "AgeRestricted": "The video is age-restricted, so its captions cannot be read without signing in.",
    "NotTranslatable": "No usable captions were found.",
    "TranslationLanguageNotAvailable": "No usable captions were found.",
}


def fetch_youtube_transcript(video_id: str, languages=("en", "tr")) -> str:
    """Fetch a transcript, preferring `languages`, then any language (human-made first).

    Raises TranscriptUnavailable when the video has no usable captions, and a plain
    RuntimeError for temporary failures that a later retry may fix.
    """
    try:
        from youtube_transcript_api import YouTubeTranscriptApi
    except ImportError as exc:
        raise RuntimeError("Install youtube-transcript-api to ingest YouTube transcripts") from exc
    api = YouTubeTranscriptApi()
    last_err = None
    for lang in languages:
        try:
            transcript = api.fetch(video_id, languages=[lang])
            return " ".join(sn.text for sn in transcript)
        except Exception as exc:
            last_err = exc
            if type(exc).__name__ != "NoTranscriptFound":
                break  # disabled, unavailable or blocked: other languages fail the same way
    if last_err is not None and type(last_err).__name__ == "NoTranscriptFound":
        try:
            available = sorted(api.list(video_id), key=lambda item: bool(getattr(item, "is_generated", True)))
            if available:
                return " ".join(sn.text for sn in available[0].fetch())
        except Exception as exc:
            last_err = exc
    message = f"Transcript unavailable for {video_id}: {last_err}"
    reason = _UNAVAILABLE_TRANSCRIPT_MESSAGES.get(type(last_err).__name__)
    if reason:
        raise TranscriptUnavailable(message, reason) from last_err
    raise RuntimeError(message) from last_err
