"""Sections: the parts of a source, taken from the source's own structure.

A section is a run of consecutive leaf chunks with a title: a playlist video, a PDF
chapter (from its bookmarks, else chapter headings on at least two pages, else page
ranges), a run of slides sharing a title, or a heading in a text document. Text with no
structure is split into parts of about `settings.section_part_chunks` chunks, named
after the concepts they teach. No AI is involved, so sections are simply rebuilt
whenever a source changes.
"""
from __future__ import annotations

import re
from pathlib import Path

from sqlalchemy import delete, func, select, update
from sqlalchemy.orm import Session

from ..config import settings
from ..models import Chunk, Concept, Document, Evidence, Section, Source

# A closing run of '#' needs a space before it, so a title such as "LINQ in C#" keeps its '#'.
HEADING_LINE = re.compile(r"^\s{0,3}(#{1,3})\s+(.+?)(?:\s+#+)?\s*$")
# "Chapter 3", "Lecture IV: ..." - the number must be followed by a heading shape, not a
# sentence ("Section 4 shows that ..." is prose, "Section 4 Results" is a heading).
CHAPTER_START = re.compile(
    r"^\W{0,3}((?i:chapter|section|part|lecture|unit|module|week)\s+(?:\d+|[IVXLCivxlc]+)\b)"
    r"(?=\s*[:.\-–—]|\s*$|\s+[A-Z0-9\"“(])"
)
CONTINUED = re.compile(r"\s*[-–—:,]?\s*[(\[]?\s*\b(?:cont(?:inued|['’]d|\.)?|contd)\s*[)\]]?\.?\s*$", re.IGNORECASE)
TITLE_EDGES = " -:,–—"
FENCE = re.compile(r"```|~~~")
HEADING_FORMATS = {".md", ".markdown", ".docx"}


def _clean_title(value: str, limit: int = 200) -> str:
    title = re.sub(r"\s+", " ", str(value or "")).strip(TITLE_EDGES)
    if len(title) <= limit:
        return title
    return title[:limit].rsplit(" ", 1)[0].rstrip(TITLE_EDGES) or title[:limit]


def _number(doc: Document, label: str) -> int | None:
    match = re.search(rf"{label}\s+(\d+)", doc.section or "", re.IGNORECASE) or re.search(rf":{label[0]}\w*?(\d+)$", doc.external_id or "")
    return int(match.group(1)) if match else None


def _documents_with_chunks(db: Session, source_id: int) -> list[tuple[Document, list[Chunk]]]:
    documents = db.execute(select(Document).where(Document.source_id == source_id).order_by(Document.id)).scalars().all()
    if not documents:
        return []
    chunks: dict[int, list[Chunk]] = {}
    for chunk in db.execute(
        select(Chunk).where(Chunk.document_id.in_([d.id for d in documents]), Chunk.chunk_kind == "leaf").order_by(Chunk.document_id, Chunk.chunk_index)
    ).scalars().all():
        chunks.setdefault(chunk.document_id, []).append(chunk)
    return [(doc, chunks[doc.id]) for doc in documents if chunks.get(doc.id)]


# --------------------------------------------------------------------------- PDF

def _outline_entries(reader, level) -> list[tuple[int, str]]:
    entries = []
    for entry in level or []:
        if isinstance(entry, list):
            continue
        try:
            page = reader.get_destination_page_number(entry) + 1
        except Exception:
            continue
        title = _clean_title(getattr(entry, "title", "") or "")
        if title and page > 0:
            entries.append((page, title))
    return entries


def _pdf_outline(path: str | None) -> list[tuple[int, str]]:
    """Chapter starts as (1-based page, title) from the PDF bookmarks, or [] when unusable.

    Uses the top level, unless it has few distinct pages (a book-title bookmark, or a
    handful of parts) and the level below has more: then the second level is used,
    keeping top-level entries without children such as "Index".
    """
    if not path or not Path(path).exists():
        return []
    try:
        from pypdf import PdfReader

        reader = PdfReader(path)
        outline = reader.outline or []
        top = _outline_entries(reader, outline)
        second = [item for level in outline if isinstance(level, list) for item in _outline_entries(reader, level)]
        entries = top
        if len({page for page, _ in top}) < 4 and len({page for page, _ in second}) > len({page for page, _ in top}):
            by_page: dict[int, str] = {page: title for page, title in top}
            by_page.update({page: title for page, title in second})  # the finer entry wins on a shared page
            entries = list(by_page.items())
        entries = sorted(dict(entries).items())
        return entries if len(entries) >= 2 else []
    except Exception:
        return []


def _pdf_first_lines(path: str | None) -> dict[int, list[str]]:
    """The first two text lines of every page, read from the PDF itself (chunking joins lines)."""
    if not path or not Path(path).exists():
        return {}
    try:
        from pypdf import PdfReader

        return {number: [line.strip() for line in (page.extract_text() or "").splitlines() if line.strip()][:2]
                for number, page in enumerate(PdfReader(path).pages, start=1)}
    except Exception:
        return {}


def _chapter_title(lines: list[str] | None, chunk_text: str) -> str | None:
    if lines is not None:
        if not lines or not CHAPTER_START.match(lines[0]):
            return None
        title = lines[0]
        if len(title) <= 20 and len(lines) > 1 and len(lines[1]) <= 80:
            title = f"{title.rstrip(':.')}: {lines[1]}"  # "Chapter 1" / "Linear Regression"
        return _clean_title(title, 120)
    match = CHAPTER_START.match(chunk_text)
    return _clean_title(chunk_text[match.start(1):match.start(1) + 80], 60) if match else None


def _pdf_segments(source: Source, groups: list[tuple[Document, list[Chunk]]]) -> list[list]:
    pages = [(_number(doc, "page") or index + 1, chunks) for index, (doc, chunks) in enumerate(groups)]
    outline = _pdf_outline(source.local_path)
    segments: list[list] = []
    if outline:
        current = None
        for page, chunks in pages:
            position = max((i for i, (start, _) in enumerate(outline) if start <= page), default=-1)
            if segments and position == current:
                segments[-1][2].extend(chunks)
            else:
                current = position
                title = outline[position][1] if position >= 0 else "Front matter"
                segments.append([title, "chapter", list(chunks)])
        return segments
    first_lines = _pdf_first_lines(source.local_path)
    headings = {page: _chapter_title(first_lines.get(page) if first_lines else None, chunks[0].text) for page, chunks in pages}
    # One match can be a title page or a stray sentence; chapter mode needs at least two.
    if sum(1 for title in headings.values() if title) >= 2:
        for page, chunks in pages:
            if headings[page] or not segments:
                segments.append([headings[page] or "Introduction", "chapter", list(chunks)])
            else:
                segments[-1][2].extend(chunks)
        return segments
    size = max(1, settings.pdf_pages_per_part)
    for start in range(0, len(pages), size):
        block = pages[start:start + size]
        first, last = block[0][0], block[-1][0]
        segments.append([f"Pages {first}–{last}" if last != first else f"Page {first}", "pages", [c for _, chunks in block for c in chunks]])
    return segments


# ------------------------------------------------------------------------ slides

def _slide_titles(path: str | None) -> dict[int, str]:
    if not path or not Path(path).exists():
        return {}
    try:
        from pptx import Presentation

        titles = {}
        for number, slide in enumerate(Presentation(path).slides, start=1):
            shape = slide.shapes.title
            if shape is not None and shape.has_text_frame and shape.text.strip():
                titles[number] = _clean_title(shape.text, 120)
        return titles
    except Exception:
        return {}


def _slide_segments(source: Source, groups: list[tuple[Document, list[Chunk]]]) -> list[list]:
    titles = _slide_titles(source.local_path)
    segments: list[list] = []
    previous_key = None
    for index, (doc, chunks) in enumerate(groups):
        number = _number(doc, "slide") or index + 1
        first_line = chunks[0].text.strip().splitlines()[0] if chunks[0].text.strip() else ""
        title = titles.get(number) or _clean_title(first_line, 120) or f"Slide {number}"
        base = _clean_title(CONTINUED.sub("", title)) or title
        key = base.casefold()
        if segments and key == previous_key:
            segments[-1][2].extend(chunks)
        else:
            segments.append([base, "slides", list(chunks)])
        previous_key = key
    return segments


# -------------------------------------------------------------------------- text

def _heading_title(line: str, strict: bool = True, capital: bool = False) -> str | None:
    match = HEADING_LINE.match(line)
    if not match:
        return None
    title = _clean_title(match.group(2), 200)
    if not title or (capital and not re.match(r"[A-Z0-9]", title)):
        return None
    if strict:
        # Chunks made before headings got their own line have the paragraph glued on;
        # such a run-on "title" reads as prose, not a heading.
        core = re.sub(r"^\d+(\.\d+)*\.?\s+", "", title)
        if len(title) > 80 or core.endswith(".") or re.search(r"[.!?]\s+\w", core):
            return None
    return title


def _overlap_length(previous: str | None, current: str) -> int:
    """How much of `current` repeats the end of `previous` (the chunk overlap).

    Compared word by word: re-splitting a long paragraph can turn a line break into a space.
    """
    if not previous or settings.chunk_overlap <= 0:
        return 0
    tokens = previous[-settings.chunk_overlap:].split()
    if not tokens:
        return 0
    match = re.match(r"\s*" + r"\s+".join(map(re.escape, tokens)), current)
    return match.end() if match else 0


def _fence_toggles(line: str) -> bool:
    """Whether a (whitespace-collapsed) line opens or closes a code fence. Balanced inline
    backtick spans such as a quoted triple backtick in prose are ignored."""
    stripped = re.sub(r"(`+)(?!`).+?(?<!`)\1(?!`)", "", line)
    return len(FENCE.findall(stripped)) % 2 == 1


def _scan_chunks(chunks: list[Chunk], strict: bool = True, capital: bool = False):
    """For each chunk: (headings in its own text, heading at its very end), outside code fences.

    The repeated overlap at the start of a chunk is skipped, so a fence or heading is seen
    once and the fence state carries across chunks. A heading on a chunk's last line has its
    text in the next chunk, so it is reported separately and opens the next chunk.
    """
    in_fence = False
    previous = None
    for chunk in chunks:
        body = chunk.text[_overlap_length(previous, chunk.text):]
        lines = [line for line in body.splitlines() if line.strip()]
        heads: list[str] = []
        last = None
        for position, line in enumerate(lines):
            if not in_fence:
                title = _heading_title(line, strict, capital)
                if title and position == len(lines) - 1 and position > 0:
                    last = title
                elif title:
                    heads.append(title)
            if _fence_toggles(line):
                in_fence = not in_fence
        previous = chunk.text
        yield chunk, heads, last


def _heading_count(groups: list[tuple[Document, list[Chunk]]], strict: bool = True, capital: bool = False) -> int:
    return sum(len(heads) + bool(last) for _, chunks in groups for _, heads, last in _scan_chunks(chunks, strict, capital))


def _text_segments(groups: list[tuple[Document, list[Chunk]]], use_headings: bool = True, strict: bool = True,
                   capital: bool = False) -> list[list]:
    segments: list[list] = []
    part = max(1, settings.section_part_chunks)
    for doc, chunks in groups:
        found: list[list] = []
        if use_headings:
            deferred = None
            for chunk, heads, last in _scan_chunks(chunks, strict, capital):
                opening = ([deferred] if deferred else []) + heads
                title = opening[0] if opening else None
                # A chunk holds one section start; a later heading in it continues into the next chunks.
                deferred = last or (opening[-1] if len(opening) > 1 else None)
                if title and not (found and found[-1][0] == title):
                    found.append([title, "heading", [chunk]])
                elif found:
                    found[-1][2].append(chunk)
                else:
                    found.append([None, "heading", [chunk]])  # text before the first heading
        if any(title for title, _, _ in found):
            for title, kind, run in found:
                if title is not None:
                    segments.append([title, kind, run])
                elif len(run) <= part:
                    segments.append([doc.title or "Introduction", "heading", run])
                else:
                    segments.extend([None, "part", run[start:start + part]] for start in range(0, len(run), part))
        elif len(chunks) <= part:
            segments.append([doc.title, "document", list(chunks)])
        else:
            # No structure: equal-sized parts, named after their concepts once those are known.
            segments.extend([None, "part", chunks[start:start + part]] for start in range(0, len(chunks), part))
    return segments


def _concept_title(db: Session, chunk_ids: list[int], number: int) -> str:
    names = db.execute(
        select(Concept.canonical_name)
        .join(Evidence, Evidence.concept_id == Concept.id)
        .where(Evidence.chunk_id.in_(chunk_ids))
        .group_by(Concept.id)
        .order_by(func.count(Evidence.id).desc(), Concept.canonical_name)
        .limit(3)
    ).scalars().all()
    return f"Part {number}: {', '.join(names)}" if names else f"Part {number}"


def build_source_sections(db: Session, source: Source) -> int:
    """Replace this source's sections with ones derived from its current structure."""
    db.execute(update(Chunk).where(Chunk.section_id.in_(select(Section.id).where(Section.source_id == source.id))).values(section_id=None))
    db.execute(delete(Section).where(Section.source_id == source.id))
    groups = _documents_with_chunks(db, source.id)
    if not groups:
        return 0
    suffix = Path(source.local_path or source.name or "").suffix.lower()
    if source.source_type == "youtube":
        segments = [[doc.title or doc.external_id or "Video", "video", list(chunks)] for doc, chunks in groups]
    elif suffix == ".pdf":
        segments = _pdf_segments(source, groups)
    elif suffix == ".pptx":
        segments = _slide_segments(source, groups)
    else:
        # Headings are structure in Markdown and Word; plain text needs at least two of them,
        # and captions or other text never use them ("# Applause #" is not a chapter).
        # Word headings are always their own paragraph; only Markdown/text can hold glued legacy titles.
        strict = suffix != ".docx"
        capital = suffix == ".txt"
        use_headings = suffix in HEADING_FORMATS or (suffix == ".txt" and _heading_count(groups, strict, capital) >= 2)
        segments = _text_segments(groups, use_headings, strict, capital)
    created = 0
    part_number = 0
    for position, (title, kind, chunks) in enumerate(segments):
        if not chunks:
            continue
        if kind == "part":
            part_number += 1
            title = _concept_title(db, [c.id for c in chunks], part_number)
        section = Section(source_id=source.id, document_id=chunks[0].document_id, position=position,
                          title=_clean_title(title or "Untitled section", 500) or "Untitled section", kind=kind, chunk_count=len(chunks))
        db.add(section)
        db.flush()
        db.execute(update(Chunk).where(Chunk.id.in_([c.id for c in chunks])).values(section_id=section.id))
        created += 1
    return created
