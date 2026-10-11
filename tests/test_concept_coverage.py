import json
import re
from contextlib import contextmanager
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

from app.config import settings
from app.models import Base, Chunk, Concept, ConceptRelation, Document, Evidence, MaintenanceJob, Section, Source
from app.services import worker as worker_module
from app.services.concepts import CONCEPT_PASS, extract_batch_concepts, pending_concept_batches
from app.services.knowledge_graph import extract_source_relations
from app.services.queue import enqueue
from app.services.sections import build_source_sections
from app.services.worker import Worker

KEYWORDS = ["Gradient Descent", "Learning Rate", "Backpropagation", "Convolution", "Attention", "Embeddings"]


class Scripted:
    """Answers concept extraction from the passages it is shown; records every call."""

    def __init__(self, errors=None):
        self.calls, self.prompts, self.errors = [], [], errors or {}

    def embed(self, texts):
        return [[1.0, 0.0, 0.0] for _ in texts]

    def generate_text(self, prompt, task="study"):
        return ""

    def generate_json(self, prompt, task="extraction", schema=None):
        self.calls.append(task)
        self.prompts.append(prompt)
        if task in self.errors:
            raise self.errors[task]
        props = (schema or {}).get("properties", {})
        if task == "extraction" and "concepts" in props:
            found = {}
            for chunk_id, text in re.findall(r"\[chunk (\d+)\] (.*)", prompt):
                for keyword in KEYWORDS:
                    if keyword.lower() in text.lower():
                        found.setdefault(keyword, []).append(int(chunk_id))
            return {"concepts": [{"name": k, "description": f"About {k}.", "aliases": [], "chunk_ids": ids} for k, ids in found.items()]}
        if task == "relationship_extraction":
            return {"relationships": []}
        return {}


@pytest.fixture()
def env(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{tmp_path / 'coverage.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    @contextmanager
    def scope():
        db = sessions()
        try:
            yield db
            db.commit()
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    from app import main
    monkeypatch.setattr(worker_module, "session_scope", scope)
    monkeypatch.setattr(main, "session_scope", scope)
    yield SimpleNamespace(sessions=sessions, client=TestClient(main.app), tmp=tmp_path)
    engine.dispose()


def _source(db, name="Course", source_type="youtube", local_path=None, docs=(("Video 1", ["Gradient descent text."]),), **extra):
    source = Source(name=name, source_type=source_type, local_path=local_path, status="done", **extra)
    db.add(source)
    db.flush()
    for index, (title, texts) in enumerate(docs):
        doc = Document(source_id=source.id, external_id=extra.get("external_ids", {}).get(index) if False else f"{name}:{index}", title=title, section=None)
        db.add(doc)
        db.flush()
        for position, text in enumerate(texts):
            db.add(Chunk(document_id=doc.id, chunk_index=position, chunk_kind="leaf", text=text, embedding_status="done"))
    db.flush()
    return source


def _chunk_ids(db, source_id):
    return db.execute(select(Chunk.id).join(Document, Chunk.document_id == Document.id).where(Document.source_id == source_id).order_by(Document.id, Chunk.chunk_index)).scalars().all()


# ------------------------------------------------------------------ batching

def test_batches_cover_every_unread_chunk_in_reading_order(env, monkeypatch):
    monkeypatch.setattr(settings, "concept_batch_chunks", 25)
    monkeypatch.setattr(settings, "concept_batch_chars", 10_000)
    with env.sessions() as db:
        source = _source(db, docs=[("A", [f"a{i} " * 10 for i in range(40)]), ("B", [f"b{i} " * 10 for i in range(20)])])
        ids = _chunk_ids(db, source.id)
        batches = pending_concept_batches(db, source.id)
        assert [cid for batch in batches for cid in batch] == ids
        assert [len(b) for b in batches] == [25, 25, 10]

        big = _source(db, name="Big", docs=[("C", ["x" * 4000 for _ in range(6)])])
        assert [len(b) for b in pending_concept_batches(db, big.id)] == [2, 2, 2]  # 10k character cap

        for chunk in db.execute(select(Chunk).where(Chunk.id.in_(ids[:30]))).scalars():
            chunk.concept_pass = CONCEPT_PASS
        db.flush()
        assert [cid for batch in pending_concept_batches(db, source.id) for cid in batch] == ids[30:]


# ---------------------------------------------------------------- extraction

def test_every_passage_is_sent_in_full_and_linked(env):
    texts = ["Gradient descent lowers the loss.", "The learning rate sets the step.", "Nothing technical here."]
    with env.sessions() as db:
        source = _source(db, docs=[("Lecture 1: Optimisation", texts)])
        ids = _chunk_ids(db, source.id)
        ai = Scripted()
        added = extract_batch_concepts(db, ai, source, ids)
        prompt = ai.prompts[0]
        assert all(text in prompt for text in texts) and "=== Lecture 1: Optimisation ===" in prompt
        assert added == 2
        names = {c.canonical_name for c in db.execute(select(Concept)).scalars()}
        assert names == {"Gradient Descent", "Learning Rate"}
        chunks = db.execute(select(Chunk).where(Chunk.id.in_(ids))).scalars().all()
        assert all(c.concept_pass == CONCEPT_PASS and c.cluster_status == "done" for c in chunks)


class Fixed(Scripted):
    def __init__(self, concepts):
        super().__init__()
        self.concepts = concepts

    def generate_json(self, prompt, task="extraction", schema=None):
        self.calls.append(task)
        self.prompts.append(prompt)
        return {"concepts": self.concepts}


def test_names_merge_ids_are_validated_and_duplicates_are_safe(env):
    with env.sessions() as db:
        existing = Concept(canonical_name="Gradient Descent", aliases_json=json.dumps(["GD"]))
        db.add(existing)
        source = _source(db, docs=[("V", ["We use GD here.", "Momentum smooths updates.", "Unrelated."])])
        ids = _chunk_ids(db, source.id)
        ai = Fixed([
            {"name": "gradient  descent", "description": "d", "chunk_ids": [ids[0], 999999]},  # spacing/case + invalid id
            {"name": "GD", "description": "d", "aliases": [], "chunk_ids": [ids[0]]},              # alias, same link twice
            {"name": "Momentum", "description": "Smooths updates", "chunk_ids": []},             # falls back to mentions
            {"name": "Quantum Gravity", "description": "x", "chunk_ids": [12345]},               # cannot be attributed
            {"name": "", "description": "x", "chunk_ids": [ids[0]]},
        ])
        extract_batch_concepts(db, ai, source, ids)
        db.flush()
        concepts = {c.canonical_name: c for c in db.execute(select(Concept)).scalars()}
        assert set(concepts) == {"Gradient Descent", "Momentum"}
        links = db.execute(select(Evidence.concept_id, Evidence.chunk_id)).all()
        assert sorted(links) == sorted([(existing.id, ids[0]), (concepts["Momentum"].id, ids[1])])


def test_failures_leave_the_batch_unread(env):
    with env.sessions() as db:
        source = _source(db, docs=[("V", ["Gradient descent."])])
        ids = _chunk_ids(db, source.id)
        assert extract_batch_concepts(db, Scripted(errors={"extraction": ValueError("bad json")}), source, ids) == 0
        assert extract_batch_concepts(db, Fixed("not a list"), source, ids) == 0
        assert db.get(Chunk, ids[0]).concept_pass == 0
        with pytest.raises(RuntimeError):
            extract_batch_concepts(db, Scripted(errors={"extraction": RuntimeError("HTTP 429")}), source, ids)


def test_later_batches_are_offered_existing_names(env):
    with env.sessions() as db:
        source = _source(db, docs=[("V", ["Gradient descent.", "More gradient descent."])])
        ids = _chunk_ids(db, source.id)
        ai = Scripted()
        extract_batch_concepts(db, ai, source, ids[:1])
        extract_batch_concepts(db, ai, source, ids[1:])
        assert '"Gradient Descent"' in ai.prompts[1].split("reuse it exactly:")[1].splitlines()[0]


# ------------------------------------------------------------------ sections

def _pdf(path, pages, outline):
    from pypdf import PdfWriter
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    for item in outline:
        if isinstance(item, tuple) and len(item) == 3:
            title, page, children = item
            parent = writer.add_outline_item(title, page)
            for child_title, child_page in children:
                writer.add_outline_item(child_title, child_page, parent=parent)
        else:
            writer.add_outline_item(*item)
    with open(path, "wb") as handle:
        writer.write(handle)


def _pdf_source(db, path, page_texts):
    source = Source(name=path.name, source_type="file", local_path=str(path), status="done")
    db.add(source)
    db.flush()
    for number, text in enumerate(page_texts, 1):
        doc = Document(source_id=source.id, external_id=f"{path.name}:p{number}", title=path.stem, section=f"page {number}")
        db.add(doc)
        db.flush()
        db.add(Chunk(document_id=doc.id, chunk_index=0, chunk_kind="leaf", text=text))
    db.flush()
    return source


def _sections(db, source_id):
    rows = db.execute(select(Section).where(Section.source_id == source_id).order_by(Section.position)).scalars().all()
    return [(s.title, s.kind, s.chunk_count) for s in rows]


def test_pdf_chapters_come_from_bookmarks(env):
    path = env.tmp / "book.pdf"
    _pdf(path, 6, [("Chapter 1: Basics", 0), ("Chapter 2: Training", 3)])
    with env.sessions() as db:
        source = _pdf_source(db, path, [f"page {i} text" for i in range(1, 7)])
        build_source_sections(db, source)
        assert _sections(db, source.id) == [("Chapter 1: Basics", "chapter", 3), ("Chapter 2: Training", "chapter", 3)]
        chunk = db.execute(select(Chunk).join(Document).where(Document.section == "page 5")).scalar_one()
        assert db.get(Section, chunk.section_id).title == "Chapter 2: Training"


def test_a_single_book_bookmark_uses_its_chapters(env):
    path = env.tmp / "book2.pdf"
    _pdf(path, 4, [("Deep Learning (book)", 0, [("Part I", 0), ("Part II", 2)])])
    with env.sessions() as db:
        source = _pdf_source(db, path, ["a", "b", "c", "d"])
        build_source_sections(db, source)
        assert [t for t, _, _ in _sections(db, source.id)] == ["Part I", "Part II"]


def test_pdf_without_bookmarks_uses_chapter_headings_then_page_ranges(env, monkeypatch):
    monkeypatch.setattr(settings, "pdf_pages_per_part", 10)
    with env.sessions() as db:
        headed = _pdf_source(db, env.tmp / "notes.pdf", ["Preface text", "Chapter 1 Linear models and more", "body", "CHAPTER 2 Neural networks", "body"])
        build_source_sections(db, headed)
        assert _sections(db, headed.id) == [("Introduction", "chapter", 1), ("Chapter 1 Linear models and more", "chapter", 2), ("CHAPTER 2 Neural networks", "chapter", 2)]

        plain = _pdf_source(db, env.tmp / "scan.pdf", [f"text {i}" for i in range(25)])
        build_source_sections(db, plain)
        assert [t for t, _, _ in _sections(db, plain.id)] == ["Pages 1–10", "Pages 11–20", "Pages 21–25"]


def test_slides_with_the_same_title_form_one_section(env):
    Presentation = pytest.importorskip("pptx").Presentation
    deck = Presentation()
    for title in ["Intro", "Backpropagation", "Backpropagation (cont.)", "Summary"]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
    path = env.tmp / "deck.pptx"
    deck.save(path)
    from app.services.ingestion import parse_file
    raws = parse_file(path)
    with env.sessions() as db:
        source = Source(name="deck.pptx", source_type="file", local_path=str(path), status="done")
        db.add(source)
        db.flush()
        for raw in raws:
            doc = Document(source_id=source.id, external_id=raw.external_id, title=raw.title, section=raw.section)
            db.add(doc)
            db.flush()
            db.add(Chunk(document_id=doc.id, chunk_index=0, chunk_kind="leaf", text=raw.text))
        db.flush()
        build_source_sections(db, source)
        assert _sections(db, source.id) == [("Intro", "slides", 1), ("Backpropagation", "slides", 2), ("Summary", "slides", 1)]


def test_text_headings_parts_and_short_documents(env, monkeypatch):
    monkeypatch.setattr(settings, "section_part_chunks", 12)
    with env.sessions() as db:
        notes = _source(db, name="notes.md", source_type="file", docs=[("notes", ["Opening words.", "# Gradient descent\nIt lowers loss.", "More detail.", "## Learning rate\nStep size."])])
        build_source_sections(db, notes)
        assert _sections(db, notes.id) == [("notes", "heading", 1), ("Gradient descent", "heading", 2), ("Learning rate", "heading", 1)]

        short = _source(db, name="short", source_type="transcript", docs=[("Short talk", ["a", "b"])])
        build_source_sections(db, short)
        assert _sections(db, short.id) == [("Short talk", "document", 2)]

        texts = [("Attention is used here." if i < 12 else "Embeddings appear here.") for i in range(30)]
        long = _source(db, name="long", source_type="transcript", docs=[("Long talk", texts)])
        long_ids = _chunk_ids(db, long.id)
        for start in range(0, len(long_ids), 10):
            extract_batch_concepts(db, Scripted(), long, long_ids[start:start + 10])
        build_source_sections(db, long)
        assert _sections(db, long.id) == [("Part 1: Attention", "part", 12), ("Part 2: Embeddings", "part", 12), ("Part 3: Embeddings", "part", 6)]

        build_source_sections(db, long)  # rebuilding replaces, never duplicates
        assert len(_sections(db, long.id)) == 3
        assert db.scalar(select(func.count()).select_from(Chunk).join(Document).where(Document.source_id == long.id, Chunk.section_id.is_(None))) == 0


def test_word_headings_are_kept(tmp_path):
    from docx import Document as WordDocument
    from app.services.ingestion import parse_file
    document = WordDocument()
    document.add_heading("Optimisation", level=1)
    document.add_paragraph("Gradient descent lowers the loss.")
    document.add_heading("Step size", level=2)
    path = tmp_path / "notes.docx"
    document.save(path)
    text = parse_file(path)[0].text
    assert "# Optimisation" in text and "## Step size" in text and "Gradient descent lowers the loss." in text


# ----------------------------------------------------------------- relations

def test_relations_cover_all_concepts_in_reading_order(env):
    with env.sessions() as db:
        texts = [f"text {i}" for i in range(90)]
        source = _source(db, docs=[("V", texts)])
        ids = _chunk_ids(db, source.id)
        for i, chunk_id in enumerate(ids):
            concept = Concept(canonical_name=f"Z{89 - i:02d} concept")  # alphabetical order is the reverse of reading order
            db.add(concept)
            db.flush()
            db.add(Evidence(concept_id=concept.id, chunk_id=chunk_id, score=1.0))
        db.flush()
        ai = Scripted()
        extract_source_relations(db, ai, source.id)
        relation_prompts = [p for p, t in zip(ai.prompts, ai.calls) if t == "relationship_extraction"]
        assert len(relation_prompts) == 3
        assert "Z89 concept" in relation_prompts[0] and "Z00 concept" in relation_prompts[2]
        assert all(p.count("- Z") <= 40 for p in relation_prompts)


# -------------------------------------------------------------- integration

def test_every_video_of_a_playlist_gets_concepts_and_a_section(env, monkeypatch):
    videos = [{"id": f"v{i}", "title": f"Lecture {i}", "channel": "", "upload_date": None, "duration_seconds": None} for i in range(1, 4)]
    topics = {"v1": "Gradient descent", "v2": "Convolution", "v3": "Attention"}
    monkeypatch.setattr(worker_module, "list_youtube_entries", lambda url: ({"kind": "playlist", "title": "Deep Learning"}, videos))
    monkeypatch.setattr(worker_module, "fetch_youtube_transcript", lambda vid: "\n\n".join(f"{topics[vid]} is explained in part {n}. " * 20 for n in range(4)))
    ai = Scripted()
    monkeypatch.setattr(worker_module, "provider", lambda: ai)
    with env.sessions() as db:
        source = Source(name="Deep Learning", source_type="youtube", uri="https://www.youtube.com/playlist?list=PLx", status="queued")
        db.add(source)
        db.flush()
        job = enqueue(db, source.id)
        db.commit()
        source_id, job_id = source.id, job.id
    Worker().process_job(job_id)
    with env.sessions() as db:
        unread = db.scalar(select(func.count()).select_from(Chunk).join(Document).where(
            Document.source_id == source_id, Chunk.chunk_kind == "leaf", Chunk.concept_pass < CONCEPT_PASS))
        assert unread == 0
        covered = db.execute(
            select(Document.title, func.count(func.distinct(Evidence.concept_id))).join(Chunk, Chunk.document_id == Document.id)
            .join(Evidence, Evidence.chunk_id == Chunk.id).where(Document.source_id == source_id).group_by(Document.title)
        ).all()
        assert dict(covered) == {"Lecture 1": 1, "Lecture 2": 1, "Lecture 3": 1}
        assert [t for t, k, _ in _sections(db, source_id)] == ["Lecture 1", "Lecture 2", "Lecture 3"]
    listed = env.client.get("/sources").json()[0]
    assert listed["section_count"] == 3 and listed["concept_count"] == 3


def test_organize_existing_backfills_older_sources(env, monkeypatch):
    with env.sessions() as db:
        source = _source(db, docs=[("Old lecture", ["Backpropagation computes gradients.", "Attention weights tokens."])])
        for chunk in db.execute(select(Chunk)).scalars():
            chunk.cluster_status = "done"  # what the old group-level extraction left behind
        db.commit()
        source_id = source.id
    queued = env.client.post("/maintenance/organize-existing").json()
    with env.sessions() as db:
        items = json.loads(db.get(MaintenanceJob, queued["id"]).items_json)
        assert {"kind": "concepts", "id": source_id} in items
    ai = Scripted()
    monkeypatch.setattr(worker_module, "provider", lambda: ai)
    Worker()._process_maintenance_job(queued["id"])
    with env.sessions() as db:
        assert db.get(MaintenanceJob, queued["id"]).status == "done"
        assert {c.canonical_name for c in db.execute(select(Concept)).scalars()} == {"Backpropagation", "Attention"}
        assert _sections(db, source_id) == [("Old lecture", "video", 2)]
    assert "relationship_extraction" in ai.calls
    second = env.client.post("/maintenance/organize-existing").json()
    with env.sessions() as db:
        assert not [i for i in json.loads(db.get(MaintenanceJob, second["id"]).items_json) if i.get("kind") == "concepts"]  # nothing left to backfill


def test_removing_a_source_removes_its_sections(env):
    with env.sessions() as db:
        source = _source(db, docs=[("V", ["text"])])
        build_source_sections(db, source)
        db.commit()
        source_id = source.id
    assert env.client.delete(f"/sources/{source_id}").status_code == 200
    with env.sessions() as db:
        assert db.scalar(select(func.count()).select_from(Section)) == 0
